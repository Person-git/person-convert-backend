import asyncio
import io
import json
from pathlib import Path
import sys
import time
import zipfile

import fitz
import httpx
from openpyxl import load_workbook
import pytest

import main
import conversion_worker
from conversion_worker import ConversionError, convert


@pytest.fixture
def pdfs(tmp_path):
    text = tmp_path / "text.pdf"
    outlines = tmp_path / "outlines.pdf"
    with fitz.open() as pdf:
        for index in range(6):
            page = pdf.new_page()
            page.insert_text((50, 50), f"Synthetic test page {index + 1}")
            for y in (100, 125, 150, 175):
                page.draw_line((50, y), (350, y))
            for x in (50, 200, 350):
                page.draw_line((x, 100), (x, 175))
            for row in range(3):
                page.insert_text((60, 116 + row * 25), f"Row {row + 1}")
                page.insert_text((210, 116 + row * 25), str(100 + row))
        pdf.save(text)
    with fitz.open() as pdf:
        for index in range(6):
            page = pdf.new_page()
            for x in range(50, 350, 5):
                page.draw_rect(fitz.Rect(x, 50, x + 3, 60))
        pdf.save(outlines)
    return text, outlines


def test_svg_all_pages_and_page_boundaries(pdfs, tmp_path):
    result = convert("svg", pdfs[0], tmp_path)
    with zipfile.ZipFile(tmp_path / result["filename"]) as archive:
        assert archive.namelist() == [f"page-{n:03d}.svg" for n in range(1, 7)]
        for item in archive.namelist():
            assert b"<svg" in archive.read(item)
    single = convert("svg", pdfs[0], tmp_path, page=6)
    assert single["filename"] == "converted.svg"
    for page in (0, -1, 7):
        with pytest.raises(ConversionError) as error:
            convert("svg", pdfs[0], tmp_path, page=page)
        assert error.value.status == 400


@pytest.mark.parametrize("kind", ["docx", "xlsx"])
def test_outline_only_pdf_rejected_before_expensive_libraries(pdfs, tmp_path, monkeypatch, kind):
    # Importing these libraries would prove the quick preflight did not protect us.
    monkeypatch.setitem(sys.modules, "pdf2docx", None)
    monkeypatch.setitem(sys.modules, "pdfplumber", None)
    with pytest.raises(ConversionError) as error:
        convert(kind, pdfs[1], tmp_path)
    assert error.value.status == 422
    assert "OCR" in error.value.detail


def test_text_pdf_converts_to_editable_word_and_excel(pdfs, tmp_path):
    from docx import Document

    result = convert("docx", pdfs[0], tmp_path)
    document = Document(tmp_path / result["filename"])
    assert "Synthetic test page 1" in " ".join(p.text for p in document.paragraphs)
    result = convert("xlsx", pdfs[0], tmp_path)
    workbook = load_workbook(tmp_path / result["filename"], read_only=True)
    assert len(workbook.sheetnames) == 6
    assert list(workbook.worksheets[0].values) == [("Row 1", "100"), ("Row 2", "101"), ("Row 3", "102")]
    workbook.close()


def test_invalid_and_password_protected_pdf(tmp_path):
    invalid = tmp_path / "bad.pdf"
    invalid.write_bytes(b"not a pdf")
    with pytest.raises(ConversionError) as error:
        convert("svg", invalid, tmp_path)
    assert error.value.status == 400
    encrypted = tmp_path / "encrypted.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(encrypted, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    with pytest.raises(ConversionError) as error:
        convert("svg", encrypted, tmp_path)
    assert error.value.status == 400


def test_api_worker_does_not_block_health_and_cleans_up(tmp_path, monkeypatch):
    marker = tmp_path / "started"
    worker = tmp_path / "slow_worker.py"
    worker.write_text(
        "import json,sys,time,pathlib\n"
        "p=json.loads(sys.argv[1]);d=pathlib.Path(p['directory'])\n"
        f"pathlib.Path({str(marker)!r}).write_text(str(d))\n"
        "time.sleep(0.5)\n"
        "(d/'converted.svg').write_text('<svg/>')\n"
        "print(json.dumps({'filename':'converted.svg','media_type':'image/svg+xml','pages':1}))\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(main, "WORKER_PATH", worker)
    monkeypatch.setattr(main, "conversion_slot", asyncio.Semaphore(1))

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
            task = asyncio.create_task(client.post("/convert/pdf-to-svg", files={"file": ("test.pdf", b"%PDF-test", "application/pdf")}))
            for _ in range(200):
                if marker.exists():
                    break
                await asyncio.sleep(0.01)
            assert marker.exists()
            started = time.perf_counter()
            health = await client.get("/health")
            assert health.json() == {"status": "ok", "version": "0.3.0"}
            assert time.perf_counter() - started < 0.2
            assert not task.done()
            response = await task
            assert response.status_code == 200
            assert "convert;dur=" in response.headers["Server-Timing"]
            assert not Path(marker.read_text()).exists()
    asyncio.run(exercise())


def test_timeout_kills_worker_and_releases_slot(tmp_path, monkeypatch):
    marker = tmp_path / "directory"
    worker = tmp_path / "hung_worker.py"
    worker.write_text(
        "import json,sys,time,pathlib\n"
        f"pathlib.Path({str(marker)!r}).write_text(json.loads(sys.argv[1])['directory'])\n"
        "time.sleep(60)\n", encoding="utf-8",
    )
    monkeypatch.setattr(main, "WORKER_PATH", worker)
    monkeypatch.setattr(main, "CONVERSION_TIMEOUT_SECONDS", 0.5)
    monkeypatch.setattr(main, "conversion_slot", asyncio.Semaphore(1))

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
            response = await client.post("/convert/pdf-to-svg", files={"file": ("test.pdf", b"%PDF-test", "application/pdf")})
            assert response.status_code == 504
            assert main.conversion_slot._value == 1
            assert marker.exists()
            assert not Path(marker.read_text()).exists()
    asyncio.run(exercise())


def test_real_api_outputs_and_validation(pdfs, monkeypatch):
    monkeypatch.setattr(main, "conversion_slot", asyncio.Semaphore(1))

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
            for extension in ("docx", "xlsx", "svg"):
                response = await client.post(f"/convert/pdf-to-{extension}", files={"file": ("test.pdf", pdfs[0].read_bytes(), "application/pdf")})
                assert response.status_code == 200, response.text[:300]
                assert float(response.headers["X-Conversion-Seconds"]) > 0
                assert response.headers["Content-Disposition"].startswith("attachment;")
                assert zipfile.is_zipfile(io.BytesIO(response.content))
            response = await client.post("/convert/pdf-to-docx", files={"file": ("outlines.pdf", pdfs[1].read_bytes(), "application/pdf")})
            assert response.status_code == 422
            assert "OCR" in response.json()["detail"]
            response = await client.post("/convert/pdf-to-svg?page=0", files={"file": ("test.pdf", pdfs[0].read_bytes(), "application/pdf")})
            assert response.status_code == 400
            response = await client.post("/convert/pdf-to-svg", files={"file": ("test.txt", b"wrong", "text/plain")})
            assert response.status_code == 400
            response = await client.post("/convert/pdf-to-svg", files={"file": ("test.pdf", b"", "application/pdf")})
            assert response.status_code == 400
    asyncio.run(exercise())


def test_compression_output_headers_and_tool_failure(pdfs, tmp_path, monkeypatch):
    import subprocess

    def fake_ghostscript(command, **kwargs):
        output = next(item.split("=", 1)[1] for item in command if item.startswith("-sOutputFile="))
        Path(output).write_bytes(pdfs[0].read_bytes())
        assert "-dPDFSETTINGS=/ebook" in command
        return subprocess.CompletedProcess(command, 0, "", "")
    monkeypatch.setattr(conversion_worker.subprocess, "run", fake_ghostscript)
    result = convert("compress", pdfs[0], tmp_path)
    assert result["filename"] == "compressed.pdf"
    assert result["headers"]["X-Original-Size"] == str(pdfs[0].stat().st_size)
    assert result["headers"]["X-Compressed-Size"] == str((tmp_path / "compressed.pdf").stat().st_size)

    def fail(command, **kwargs):
        raise subprocess.TimeoutExpired(command, 120)
    monkeypatch.setattr(conversion_worker.subprocess, "run", fail)
    with pytest.raises(ConversionError) as error:
        convert("compress", pdfs[0], tmp_path)
    assert error.value.status == 504


def test_queue_deadline_and_oversized_upload(monkeypatch):
    monkeypatch.setattr(main, "conversion_slot", asyncio.Semaphore(1))
    monkeypatch.setattr(main, "QUEUE_TIMEOUT_SECONDS", 0.05)
    monkeypatch.setattr(main, "MAX_FILE_SIZE_MB", 1)

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
            await main.conversion_slot.acquire()
            try:
                response = await client.post("/convert/pdf-to-svg", files={"file": ("test.pdf", b"%PDF-test", "application/pdf")})
                assert response.status_code == 503
                assert response.headers["Retry-After"] == "30"
                assert (await client.get("/health")).status_code == 200
            finally:
                main.conversion_slot.release()
            response = await client.post("/convert/pdf-to-svg", files={"file": ("big.pdf", b"x" * (1024 * 1024 + 1), "application/pdf")})
            assert response.status_code == 413
    asyncio.run(exercise())
