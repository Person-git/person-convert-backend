"""Smoke test only generated, non-confidential PDFs against a chosen API."""
import argparse
import io
import json
from pathlib import Path
import time
import zipfile

import fitz
import httpx
from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument("--base-url", required=True)
parser.add_argument("--output", required=True)
args = parser.parse_args()
output = Path(args.output)
output.mkdir(parents=True, exist_ok=True)
fixture = Path("tmp/synthetic-benchmark/synthetic-six-pages.pdf").read_bytes()
rows = []
with httpx.Client(base_url=args.base_url, timeout=150) as client:
    ready_until = time.monotonic() + 60
    while True:
        try:
            health = client.get("/health", timeout=10)
            health.raise_for_status()
            break
        except httpx.HTTPError:
            if time.monotonic() >= ready_until:
                raise
            time.sleep(1)
    health.raise_for_status()
    assert health.json() == {"status": "ok", "version": "0.3.0"}, health.text
    cors = client.options("/convert/pdf-to-docx", headers={"Origin": "null", "Access-Control-Request-Method": "POST"})
    assert cors.status_code == 200
    assert cors.headers["Access-Control-Allow-Origin"] == "*"
    for kind in ("docx", "xlsx", "svg", "compress"):
        started = time.perf_counter()
        endpoint = "/convert/pdf-compress" if kind == "compress" else f"/convert/pdf-to-{kind}"
        response = client.post(endpoint, files={"file": ("synthetic-six-pages.pdf", fixture, "application/pdf")})
        response.raise_for_status()
        assert float(response.headers["X-Conversion-Seconds"]) > 0
        if kind == "docx":
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                assert b"Synthetic public benchmark page 1" in archive.read("word/document.xml")
        elif kind == "xlsx":
            workbook = load_workbook(io.BytesIO(response.content), read_only=True)
            assert len(workbook.sheetnames) == 6
            workbook.close()
        elif kind == "svg":
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                assert archive.namelist() == [f"page-{i:03d}.svg" for i in range(1, 7)]
        else:
            with fitz.open(stream=response.content, filetype="pdf") as pdf:
                assert pdf.page_count == 6
                assert "Synthetic public benchmark page 1" in pdf[0].get_text()
            assert response.headers["X-Original-Size"] == str(len(fixture))
            assert response.headers["X-Compressed-Size"] == str(len(response.content))
        row = {"kind": kind, "wall_seconds": round(time.perf_counter()-started, 3),
               "server_seconds": response.headers["X-Conversion-Seconds"], "bytes": len(response.content)}
        rows.append(row)
        print(json.dumps(row), flush=True)
    # Deliberately outlined drawings with no text; never use a private source here.
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.draw_rect(fitz.Rect(50, 50, 200, 70))
        outline = pdf.tobytes()
    for kind in ("docx", "xlsx"):
        response = client.post(f"/convert/pdf-to-{kind}", files={"file": ("synthetic-outline.pdf", outline, "application/pdf")})
        assert response.status_code == 422, response.text
        assert "OCR" in response.json()["detail"]
    response = client.post("/convert/pdf-to-svg?page=0", files={"file": ("synthetic.pdf", fixture, "application/pdf")})
    assert response.status_code == 400
(output/"smoke-test.json").write_text(json.dumps(rows, indent=2), encoding="utf-8")
print("All four conversions, CORS, no-text errors and page validation passed", flush=True)
